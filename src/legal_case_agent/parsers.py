from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ParseError(RuntimeError):
    pass


@dataclass(frozen=True)
class SourceChunk:
    document_id: str
    source_label: str
    text: str


@dataclass(frozen=True)
class ParseResult:
    chunks: list[SourceChunk]
    warnings: list[str]


def parse_document(document_id: str, original_name: str, path: Path) -> ParseResult:
    extension = path.suffix.lower()
    try:
        if extension == ".txt":
            text = path.read_text(encoding="utf-8-sig")
            parts = [part.strip() for part in text.splitlines() if part.strip()]
            return ParseResult(
                [SourceChunk(document_id, f"{original_name}，第 {i} 段", part) for i, part in enumerate(parts, 1)],
                [],
            )
        if extension == ".docx":
            document = Document(path)
            parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table_index, table in enumerate(document.tables, 1):
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append(f"表格 {table_index}：\n" + "\n".join(rows))
            return ParseResult(
                [SourceChunk(document_id, f"{original_name}，第 {i} 段", part) for i, part in enumerate(parts, 1)],
                [],
            )
        if extension == ".pdf":
            chunks = []
            for i, page in enumerate(PdfReader(path).pages, 1):
                text = (page.extract_text() or "").strip()
                if text:
                    chunks.append(SourceChunk(document_id, f"{original_name}，第 {i} 页", text))
            warnings = [] if chunks else [f"{original_name}：未提取到文字，可能需要 OCR"]
            return ParseResult(chunks, warnings)
        if extension in {".png", ".jpg", ".jpeg"}:
            return ParseResult([], [f"{original_name}：需要人工查看或视觉模型识别"])
    except Exception as exc:
        raise ParseError(f"{original_name}：解析失败") from exc
    raise ParseError(f"{original_name}：不支持的文件类型")

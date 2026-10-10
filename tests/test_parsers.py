from legal_case_agent.parsers import parse_document


def test_txt_chunks_include_stable_source(tmp_path):
    path = tmp_path / "聊天记录.txt"
    path.write_text("2026-02-01 公司通知解除劳动合同", encoding="utf-8")
    result = parse_document("doc1", path.name, path)
    assert result.chunks[0].source_label == "聊天记录.txt，第 1 段"
    assert "解除劳动合同" in result.chunks[0].text


def test_image_is_marked_for_manual_review(tmp_path):
    path = tmp_path / "证据.jpg"
    path.write_bytes(b"image")
    result = parse_document("doc2", path.name, path)
    assert not result.chunks
    assert "人工查看" in result.warnings[0]


def test_docx_table_is_included(tmp_path):
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "月份"
    table.cell(0, 1).text = "工资"
    table.cell(1, 0).text = "一月"
    table.cell(1, 1).text = "8000"
    path = tmp_path / "工资.docx"
    document.save(path)

    result = parse_document("doc", path.name, path)
    assert any("月份 | 工资" in chunk.text and "一月 | 8000" in chunk.text for chunk in result.chunks)
